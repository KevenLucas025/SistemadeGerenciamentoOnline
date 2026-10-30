from io import BytesIO
import pandas as pd
from django.http import Http404, HttpResponse
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from clientes.models import Cliente
from produtos.models import Produto
from datetime import datetime
from decimal import Decimal, InvalidOperation
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import JsonResponse
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST


# LEVA ATÉ A PÁGINA DE CADASTRAR PRODUTOS EM MASSA
@login_required
def ir_pagina_produtos_massa(request):
    return render(
        request,
        "cadastrar_produtos_massa.html",
    )
    
    
@login_required
def baixar_planilha_exemplo(request, tipo):
    if tipo == "produtos":
        nome_arquivo = "Produtos Exemplo.xlsx"
        nome_aba = "Produtos"
    
        dados = {
            "Produto": [
                "Notebook Exemplo",
                "Mouse sem Fio",
                "Impressora Exemplo",
                "Monitor Full HD",
            ],
            "Quantidade": [10, 50, 5, 20],
            "Valor Unitário": [4500.00, 150.00, 1200.00, 900.00],
            "Desconto": [5, "Sem desconto", 10, "Sem desconto"],
            "Data do Cadastro": [
                "10/05/2024",
                "10/04/2024",
                "10/03/2024",
                "10/02/2024",
            ],
            "Código do Item": [
                "Gerado automaticamente",
                "Gerado automaticamente",
                "Gerado automaticamente",
                "Gerado automaticamente",
            ],
            "Cliente": [
                "João da Silva",
                "Maria Oliveira",
                "Pedro Santos",
                "Carla Lima",
            ],
            "Descrição": [
                "Notebook com 16GB RAM e SSD",
                "Mouse sem fio",
                "Impressora multifuncional",
                "Monitor de 24 polegadas",
            ],
        }
    elif tipo == "usuarios":
        nome_arquivo = "Usuarios Exemplos.xlsx"
        nome_aba = "Usuários"

        dados = {
            "Nome": [
                "Usuário Exemplo 1",
                "Usuário Exemplo 2",
                "Usuário Exemplo 3",
                "Usuário Exemplo 4",
            ],
            "Usuário": [
                "Exemplo 1",
                "Exemplo 2",
                "Exemplo 3",
                "Exemplo 4",
            ],
            "Senha": [
                "senha123",
                "senha456",
                "senha789",
                "senha101",
            ],
            "Confirmar Senha": [
                "senha123",
                "senha456",
                "senha789",
                "senha101",
            ],
            "CEP": [
                "00000-0001",
                "00000-0002",
                "00000-0003",
                "00000-0004",
            ],
            "Endereço": [
                "Rua A",
                "Rua B",
                "Rua C",
                "Rua D",
            ],
            "Número": [123, 456, 789, 101],
            "Cidade": [
                "Campinas",
                "Sumaré",
                "São Paulo",
                "Monte Mor",
            ],
            "Bairro": [
                "Centro",
                "Zona Sul",
                "Zona Norte",
                "Zona Leste",
            ],
            "Estado": ["SP", "SP", "SP", "SP"],
            "Complemento": [
                "Apto 1",
                "Opcional",
                "Opcional",
                "Casa 2",
            ],
            "Telefone": [
                "(11) 00000-0001",
                "(11) 00000-0002",
                "(11) 00000-0003",
                "(11) 00000-0004",
            ],
            "E-mail": [
                "usuario1@exemplo.com",
                "usuario2@exemplo.com",
                "usuario3@exemplo.com",
                "usuario4@exemplo.com",
            ],
            "Data de Nascimento": [
                "01/01/2000",
                "02/02/1995",
                "03/03/1990",
                "04/04/1985",
            ],
            "RG": [
                "00.000.000-1",
                "00.000.000-2",
                "00.000.000-3",
                "00.000.000-4",
            ],
            "CPF": [
                "000.000.000-01",
                "000.000.000-02",
                "000.000.000-03",
                "000.000.000-04",
            ],
            "CNPJ": [
                "00.000.000/0001-01",
                "00.000.000/0001-02",
                "00.000.000/0001-03",
                "00.000.000/0001-04",
            ],
            "Acesso": [
                "Convidado",
                "Usuário",
                "Usuário",
                "Convidado",
            ],
        }
    else:
        raise Http404("Tipo de planilha não encontrado.")
    
    df = pd.DataFrame(dados)
    
    arquivo = BytesIO()
    
    df.to_excel(
        arquivo,
        index=False,
        engine="openpyxl",
        sheet_name=nome_aba,
    )
    
    arquivo.seek(0)
    
    workbook = load_workbook(arquivo)
    worksheet = workbook[nome_aba]
    
    # ESTILO DO CABEÇALHO
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="334155",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # AJUSTA A LARGURA DAS COLUNAS
    for coluna in worksheet.columns:
        maior_largura = 0
        letra_coluna = get_column_letter(coluna[0].column)

        for cell in coluna:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            if cell.value is not None:
                maior_largura = max(
                    maior_largura,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[letra_coluna].width = min(
            maior_largura + 2,
            45,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # SALVA A FORMATAÇÃO NO ARQUIVO EM MEMÓRIA
    arquivo_formatado = BytesIO()
    workbook.save(arquivo_formatado)
    arquivo_formatado.seek(0)

    # DEVOLVE O ARQUIVO PARA DOWNLOAD
    response = HttpResponse(
        arquivo_formatado.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{nome_arquivo}"'
    )

    return response

